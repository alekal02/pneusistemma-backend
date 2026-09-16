import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

src = Path(r"c:\Users\Sistemma\Downloads\modelo-pneus (3) (4) PREENCHIDO.xlsx")
out = Path(r"c:\Users\Sistemma\Downloads\aparecida-pneus-import.json")

wb = load_workbook(src, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
hdr = list(rows[0])


def cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def fogo_str(v):
    s = cell(v)
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s


def tipo_de_obs(obs):
    o = re.sub(r"[^A-Z0-9]", "", str(obs or "").upper())
    if o.startswith("1"):
        return "direcional"
    if o.startswith("2"):
        return "borrachudo"
    return "misto"


itens = []
vistos = Counter()
descarte = set()
sn = 0

for i, r in enumerate(rows[1:], start=2):
    if all(v is None or str(v).strip() == "" for v in r):
        continue
    d = {hdr[j]: r[j] for j in range(len(hdr))}
    obs = cell(d.get(hdr[-1]))
    cond_raw = cell(d.get("condicao"))
    fogo = fogo_str(d.get("numero_fogo"))
    placeholder = (not fogo) or re.fullmatch(r"(?i)s/?n|sem\s*n(ume?ro)?", fogo)
    if placeholder:
        fogo = ""
        sn += 1
    else:
        vistos[fogo.upper()] += 1
        if vistos[fogo.upper()] > 1:
            fogo = f"{fogo}-{vistos[fogo.upper()]}"
    vida = cell(d.get("VIDA_RECAP"))
    cond_up = cond_raw.upper()
    if cond_up.startswith("DESCARTE"):
        condicao = "recapado"
        if fogo:
            descarte.add(fogo)
        elif obs:
            descarte.add(f"__linha_{i}")
    elif cond_up.startswith("REC"):
        condicao = "recapado"
    else:
        condicao = "novo"
    item = {
        "__linha": i,
        "numero_fogo": fogo,
        "tipo_pneu": tipo_de_obs(obs),
        "condicao": condicao,
        "vida_recap": vida,
        "marca": cell(d.get("MARCA")),
        "modelo": cell(d.get("modelo")),
        "medida": cell(d.get("medida")).replace(",", "."),
        "dot": cell(d.get("dot/serie")),
        "sulco_inicial_mm": cell(d.get("SULCO_INICIAL_MM")),
        "custo": cell(d.get("custo")),
        "nota_fiscal": cell(d.get("nota_fiscal")),
        "observacoes": obs if not cond_up.startswith("DESCARTE") else (obs or cond_raw),
    }
    if cond_up.startswith("DESCARTE") and cond_raw not in (item["observacoes"] or ""):
        extra = cond_raw
        item["observacoes"] = f"{item['observacoes']}; {extra}".strip("; ")
    itens.append(item)

payload = {
    "itens": itens,
    "descarte_fogos": sorted(descarte),
    "meta": {
        "linhas": len(itens),
        "sn_gerar": sn,
        "unicos": len({x["numero_fogo"] for x in itens if x["numero_fogo"]}),
        "descarte": len(descarte),
        "fogos_digitais": sum(1 for x in itens if str(x["numero_fogo"]).isdigit()),
        "amostra": [x["numero_fogo"] for x in itens[:8]],
    },
}
out.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
print(json.dumps(payload["meta"], ensure_ascii=False, indent=2))
print("wrote", out, "bytes", out.stat().st_size)
